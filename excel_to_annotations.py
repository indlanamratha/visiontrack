"""Convert the research team's coding sheet (Excel) into annotations.json for render.py.

Each spreadsheet row = one person in one time frame (e.g. "time00:15-00:20").
Rows that share a time frame become one segment; the segment's risk is the
highest "overall risk level" among its rows, and each channel score
(gesture / facial / voice) is the highest coded value.

Usage: python excel_to_annotations.py sample_output.xlsx annotations.json [--video-id clip00001]
Tracker IDs still have to be linked to people (fill "tracks" after running track.py).
"""
import argparse, json, re
from collections import OrderedDict
import openpyxl

COLS = {  # spreadsheet column letter -> field (matches the sample coding sheet layout)
    "A": "video", "D": "timeframe", "E": "person", "F": "activity", "G": "weapon",
    "H": "stance", "I": "sudden", "J": "gesture", "K": "gesture_esc", "Q": "facial_esc",
    "N": "face_visible", "Y": "speaking", "AA": "voice_esc", "AD": "sentence",
    "AG": "location", "AH": "type", "AM": "risk",
}

def secs(mmss):
    m, s = mmss.split(":")
    return int(m) * 60 + float(s)

def num(v):
    try: return int(v)
    except (TypeError, ValueError): return None

ap = argparse.ArgumentParser()
ap.add_argument("xlsx"); ap.add_argument("out")
ap.add_argument("--video-id", default=None)
a = ap.parse_args()

ws = openpyxl.load_workbook(a.xlsx, data_only=True).active
rows = []
for r in ws.iter_rows(min_row=3):  # row 1 = headers, row 2 = coding instructions
    rec = {f: ws[f"{c}{r[0].row}"].value for c, f in COLS.items()}
    if not rec["timeframe"] or (a.video_id and rec["video"] != a.video_id):
        continue
    m = re.search(r"(\d+:\d+(?:\.\d+)?)\s*-\s*(\d+:\d+(?:\.\d+)?)", str(rec["timeframe"]))
    if not m:
        continue
    rec["start"], rec["end"] = secs(m.group(1)), secs(m.group(2))
    rows.append(rec)

segs = OrderedDict()
people = OrderedDict()
for r in rows:
    key = (r["start"], r["end"])
    sg = segs.setdefault(key, {"start": r["start"], "end": r["end"], "risk": 0,
                               "channels": {"gesture": None, "facial": None, "voice": None},
                               "people": {}, "triggers": []})
    pid = str(r["person"]).strip()
    role = "police" if pid.upper().startswith("P") else "suspect" if pid.upper().startswith("S") else "bystander"
    people.setdefault(pid, {"role": role, "label": {"police": "Officer", "suspect": "Suspect"}.get(role, "Person"), "tracks": []})
    weapon = (r["weapon"] or "none").strip().lower()
    st = {"activity": r["activity"] or "", "weapon": weapon, "sudden": num(r["sudden"]) or 0}
    if r["gesture"]: st["gesture"] = str(r["gesture"]).lower()
    sg["people"][pid] = st
    sg["risk"] = max(sg["risk"], num(r["risk"]) or 0)
    for ch, col in (("gesture", "gesture_esc"), ("facial", "facial_esc"), ("voice", "voice_esc")):
        v = num(r[col])
        if v is not None:
            sg["channels"][ch] = max(sg["channels"][ch] or 0, v)
    if weapon not in ("none", ""): sg["triggers"].append(f"{pid}: {weapon} visible" + (f" ({r['gesture']})" if r["gesture"] else ""))
    if num(r["stance"]): sg["triggers"].append(f"{pid}: fighting stance")
    if num(r["sudden"]): sg["triggers"].append(f"{pid}: sudden movement")
    if r["sentence"]:
        sg["triggers"].append(f"{pid} yelling: \"{r['sentence']}\"")
        sg["speech"] = {"who": pid, "text": r["sentence"]}
for sg in segs.values():
    if sg["risk"] >= 4:
        sg["headline"] = " · ".join(t.split(": ", 1)[-1] for t in sg["triggers"][:2])

first = rows[0] if rows else {}
out = {"clip": {"id": first.get("video", ""), "location": first.get("location", ""), "type": first.get("type", ""),
                "offset": 0.0, "start": min((s["start"] for s in segs.values()), default=0),
                "end": max((s["end"] for s in segs.values()), default=0)},
       "people": people, "segments": list(segs.values())}
json.dump(out, open(a.out, "w"), indent=1)
print(f"{len(rows)} rows -> {len(segs)} segments, people: {', '.join(people)}")
